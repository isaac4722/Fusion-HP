#!/usr/bin/env python3
# ============================================================================
#  LuminaPresentation / LuminaPresentation Suite - Copyright (c) 2026 Isaac. Licencia View-Only.
# ============================================================================
#  validate_pptx.py : GATE de validación EXTERNA de los exportadores v5.0.0.
#
#    * PPTX → python-pptx: abre el paquete OPC (valida estructura + rels +
#      content types), cuenta slides, mide el lienzo en EMU y extrae el texto
#      real de las diapositivas (a:t). Si PowerPoint-reality falla, falla el CI.
#    * PDF  → pypdf: abre el documento (valida xref/objetos), cuenta páginas
#      y extrae texto (valida WinAnsi + base-14).
#
#  Uso:  python tools/validate_pptx.py <muestra.pptx> <muestra.pdf>
#  Exit 0 = OK · distinto de 0 = gate rojo.
# ============================================================================
import sys


def fail(msg):
    print("VALIDATE-EXPORT FAIL:", msg)
    sys.exit(1)


def main():
    if len(sys.argv) < 3:
        fail("uso: validate_pptx.py <muestra.pptx> <muestra.pdf>")
    pptx_path, pdf_path = sys.argv[1], sys.argv[2]

    # ------------------------------------------------------------- PPTX
    try:
        from pptx import Presentation
        from pptx.util import Emu
    except ImportError:
        fail("python-pptx no está instalado (pip install python-pptx)")

    try:
        prs = Presentation(pptx_path)
    except Exception as ex:
        fail("python-pptx no pudo abrir el PPTX: %s" % ex)

    n = len(prs.slides)
    if n != 3:
        fail("se esperaban 3 diapositivas, hay %d" % n)

    if prs.slide_width != 12192000 or prs.slide_height != 6858000:
        fail("tamaño de diapositiva incorrecto: %dx%d EMU (se esperaba 12192000x6858000)"
             % (prs.slide_width, prs.slide_height))

    def slide_text(idx):
        texts = []
        for shape in prs.slides[idx].shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    for run in para.runs:
                        texts.append(run.text)
        return " ".join(texts)

    s0 = slide_text(0)
    if "Culto de prueba" not in s0:
        fail("la diapositiva 1 no contiene el título: %r" % s0[:120])
    s1 = slide_text(1)
    if "Juan 3:16" not in s1:
        fail("la diapositiva 2 no contiene la referencia: %r" % s1[:120])
    if "amó" not in s1:
        fail("la diapositiva 2 no conserva el acento (amó): %r" % s1[:120])

    print("PPTX OK: %d diapositivas, 12192000x6858000 EMU, texto y acentos correctos" % n)

    # -------------------------------------------------------------- PDF
    try:
        from pypdf import PdfReader
    except ImportError:
        try:
            from PyPDF2 import PdfReader  # variante antigua del nombre
        except ImportError:
            fail("pypdf no está instalado (pip install pypdf)")

    try:
        reader = PdfReader(pdf_path)
    except Exception as ex:
        fail("pypdf no pudo abrir el PDF: %s" % ex)

    npages = len(reader.pages)
    if npages != 3:
        fail("se esperaban 3 páginas en el PDF, hay %d" % npages)

    box = reader.pages[0].mediabox
    if abs(float(box.width) - 960.0) > 0.5 or abs(float(box.height) - 540.0) > 0.5:
        fail("MediaBox incorrecto: %sx%s (se esperaba 960x540 pt)" % (box.width, box.height))

    try:
        t0 = reader.pages[0].extract_text() or ""
        t1 = reader.pages[1].extract_text() or ""
    except Exception as ex:
        fail("extracción de texto del PDF falló: %s" % ex)

    if "Culto de prueba" not in t0:
        fail("página 1 del PDF sin el título: %r" % t0[:120])
    if "Juan 3:16" not in t1:
        fail("página 2 del PDF sin la referencia: %r" % t1[:120])
    if "amó" not in t1:
        fail("página 2 del PDF sin el acento (amó): %r" % t1[:120])

    print("PDF OK: %d páginas, 960x540 pt, texto y acentos correctos" % npages)
    print("VALIDATE-EXPORT PASS")
    sys.exit(0)


if __name__ == "__main__":
    main()
